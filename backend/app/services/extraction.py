import re
from pathlib import Path
from typing import Any

# pyrefly: ignore [missing-import]
import fitz
import pandas as pd

from docx import Document as WordDocument


def extract_pages(path: Path) -> list[tuple[int | None, str]]:
    """Extract page number and text from PDF, DOCX, XLSX, or TXT documents."""
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        pages = []
        with fitz.open(path) as pdf:
            for i, page in enumerate(pdf):
                text = page.get_text("text").strip()
                if text:
                    pages.append((i + 1, text))
        return pages or [(1, "")]

    if suffix == ".docx":
        doc = WordDocument(path)
        parts = []
        for p in doc.paragraphs:
            if p.text.strip():
                parts.append(p.text.strip())
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    parts.append(row_text)
        full_text = "\n\n".join(parts)
        return [(None, full_text)]

    if suffix == ".xlsx":
        sheets = pd.read_excel(path, sheet_name=None)
        pages = []
        for i, (sheet_name, frame) in enumerate(sheets.items(), 1):
            if frame.empty:
                continue
            # Format sheet cleanly as text
            csv_text = frame.dropna(how="all").to_csv(index=False)
            pages.append((i, f"Sheet: {sheet_name}\n{csv_text}"))
        return pages or [(None, "")]

    if suffix == ".txt":
        content = path.read_text(encoding="utf-8", errors="replace").strip()
        return [(None, content)]

    raise ValueError(f"Unsupported file format: {suffix}. Supported: .pdf, .docx, .xlsx, .txt")


def chunk_pages(
    pages: list[tuple[int | None, str]],
    chunk_size: int = 1000,
    overlap: int = 200,
) -> list[dict[str, Any]]:
    """
    Split text into overlapping chunks with smart boundary awareness (sentences/paragraphs).
    """
    chunks = []
    chunk_idx = 0

    for page, raw_text in pages:
        if not raw_text.strip():
            continue

        # Split into paragraph blocks
        paragraphs = [p.strip() for p in raw_text.split("\n") if p.strip()]
        normalized_text = "\n\n".join(paragraphs)

        start = 0
        text_len = len(normalized_text)

        while start < text_len:
            end = min(start + chunk_size, text_len)
            
            if end < text_len:
                # Try to break at paragraph boundary, sentence boundary, or space boundary
                p_bound = normalized_text.rfind("\n\n", start, end)
                if p_bound > start + (chunk_size // 2):
                    end = p_bound + 2
                else:
                    s_match = re.search(r"[.!?]\s", normalized_text[start:end])
                    if s_match:
                        # Find last sentence end in the segment
                        matches = [m.end() for m in re.finditer(r"[.!?]\s", normalized_text[start:end])]
                        if matches and (matches[-1] > chunk_size // 3):
                            end = start + matches[-1]
                    else:
                        space_bound = normalized_text.rfind(" ", start, end)
                        if space_bound > start + (chunk_size // 3):
                            end = space_bound + 1

            chunk_text = normalized_text[start:end].strip()
            if chunk_text:
                chunks.append({
                    "chunk_id": chunk_idx,
                    "page": page,
                    "text": chunk_text,
                })
                chunk_idx += 1

            step = end - start
            if step <= overlap:
                start = end
            else:
                start = end - overlap

    return chunks
